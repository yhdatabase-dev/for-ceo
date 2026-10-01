"""docx 생성 — 원본 `cgr/docx_export.py:text_to_docx` 이관.

원본 로직 요약
- 텍스트를 문단 단위로 python-docx Document 에 write
- 제목·부제·푸터 삽입
- 바이트 반환

TODO(이관): 원본 로직 이관. 지금은 최소 구현.
"""
from __future__ import annotations

import io

from docx import Document


def text_to_docx(
    text: str,
    *,
    title: str = "",
    subtitle: str = "",
    footer_note: str = "",
) -> bytes:
    """텍스트 → .docx 바이트."""
    doc = Document()
    if title:
        doc.add_heading(title, level=0)
    if subtitle:
        doc.add_paragraph(subtitle)
    for line in text.split("\n"):
        doc.add_paragraph(line)
    if footer_note:
        doc.add_paragraph()
        doc.add_paragraph(footer_note)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def wr_comparison_to_docx(rows, effective_date: str = "") -> bytes:
    """조별 대비표 .docx — 원본 `wr_comparison_to_docx` 이관.

    TODO: 원본 로직 이관 (테이블 삽입).
    """
    doc = Document()
    doc.add_heading("취업규칙 대비표", level=0)
    if effective_date:
        doc.add_paragraph(f"시행일: {effective_date}")
    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = "조번호"
    hdr[1].text = "원문"
    hdr[2].text = "수정본"
    hdr[3].text = "근거"
    for r in rows:
        row = table.add_row().cells
        row[0].text = str(r.get("조번호", ""))
        row[1].text = str(r.get("원문", ""))
        row[2].text = str(r.get("수정본", ""))
        row[3].text = str(r.get("근거", ""))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
