"""docx 파서 — python-docx 사용.

원본: `cgr/parsers/docx.py:parse_docx(path) -> str` 를 그대로 이관.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document


def parse(path: Path) -> str:
    """docx → 텍스트. 문단 + 테이블 셀 순회."""
    doc = Document(str(path))
    lines: list[str] = []
    for p in doc.paragraphs:
        text = (p.text or "").strip()
        if text:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = (cell.text or "").strip()
                if text:
                    lines.append(text)
    return "\n".join(lines)
